import csv
import io
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def _flatten_value(value):
    if value is None:
        return ''
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_csv(filename, columns, rows, summary=None):
    """Generate a CSV file response."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    if summary:
        writer.writerow(['Summary'])
        for key, val in summary.items():
            writer.writerow([key.replace('_', ' ').title(), val])
        writer.writerow([])

    writer.writerow([col['label'] for col in columns])
    for row in rows:
        writer.writerow([_flatten_value(row.get(col['key'], '')) for col in columns])

    response = HttpResponse(buffer.getvalue(), content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return response


def export_xlsx(filename, columns, rows, summary=None, sheet_name='Report'):
    """Generate an Excel file response."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    row_idx = 1
    if summary:
        ws.cell(row=row_idx, column=1, value='Summary').font = Font(bold=True)
        row_idx += 1
        for key, val in summary.items():
            ws.cell(row=row_idx, column=1, value=key.replace('_', ' ').title())
            ws.cell(row=row_idx, column=2, value=_flatten_value(val))
            row_idx += 1
        row_idx += 1

    for col_idx, col in enumerate(columns, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=col['label'])
        cell.font = Font(bold=True)
    row_idx += 1

    for row in rows:
        for col_idx, col in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=_flatten_value(row.get(col['key'], '')))
        row_idx += 1

    for col_idx in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 18

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    return response


def build_export_response(format_type, filename, columns, rows, summary=None):
    if format_type == 'xlsx':
        return export_xlsx(filename, columns, rows, summary)
    return export_csv(filename, columns, rows, summary)
